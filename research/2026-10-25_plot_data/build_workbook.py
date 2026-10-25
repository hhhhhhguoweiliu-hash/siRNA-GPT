from pathlib import Path
import csv, json, statistics
from decimal import Decimal
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

d=Path(__file__).resolve().parent
rows=list(csv.DictReader((d/'source.tsv').open(encoding='utf-8-sig'),delimiter='\t'))
assert len(rows)==8 and len({r['sample'] for r in rows})==8
nor=sorted([r for r in rows if r['state']=='Normoxia'],key=lambda r:r['sample'])
hyp=sorted([r for r in rows if r['state']=='Hypoxia'],key=lambda r:r['sample'])
assert len(nor)==len(hyp)==4
means={name:sum(Decimal(r['TPM']) for r in vals)/4 for name,vals in [('Normoxia',nor),('Hypoxia',hyp)]}
assert means['Normoxia']==Decimal('20.09929925') and means['Hypoxia']==Decimal('102.390068')
wb=Workbook();s=wb.active;s.title='Prism_columns'
s.append(['Normoxia','Hypoxia'])
for a,b in zip(nor,hyp):s.append([float(a['TPM']),float(b['TPM'])])
for c in s[1]:c.comment=Comment('Each column contains four dishes from the same culture batch. Row alignment is for layout only and does not imply pairing. No siRNA or transfection reagent.','Data provenance')
s.freeze_panes='A2';s.column_dimensions['A'].width=24;s.column_dimensions['B'].width=24
t=wb.create_sheet('Samples_long');t.append(['Sample','Condition','VEGFA_TPM'])
for r in nor+hyp:t.append([r['sample'],r['state'],float(r['TPM'])])
t.freeze_panes='A2';t.auto_filter.ref=t.dimensions
for key,width in [('A',18),('B',22),('C',22)]:t.column_dimensions[key].width=width
u=wb.create_sheet('Summary');u.append(['Condition','Dish_count','Mean_VEGFA_TPM'])
for name in ['Normoxia','Hypoxia']:u.append([name,4,float(means[name])])
u.append([]);u.append(['Ratio_Hypoxia_over_Normoxia',None,float(means['Hypoxia']/means['Normoxia'])])
for key,width in [('A',38),('B',18),('C',24)]:u.column_dimensions[key].width=width
n=wb.create_sheet('说明')
notes=[('内容','截图所用的8个真实VEGFA TPM原值；不是从图上估读。'),
('作图方式','Prism_columns：两列各4个点。列间行对齐仅为排版，不代表配对，按两组非配对数据处理。'),
('样本','同一培养批次；每状态4皿，分别提取RNA；不是4轮独立培养。'),
('缺氧条件','Hyp1–4：1% O2、连续24小时；Nor1–4同期常氧培养。'),
('干预','全部8皿均无siRNA、无转染试剂，为未干预基线。'),
('图中横线','显示组均值；此表未新增显著性检验或误差线。'),
('原表','research/2026-10-20_integrity_arial/paper/data/VEGFA_eight_samples_TPM.tsv'),
('GitHub来源','https://github.com/hhhhhhguoweiliu-hash/siRNA-GPT/blob/main/research/2026-10-20_integrity_arial/paper/data/VEGFA_eight_samples_TPM.tsv'),
('Git blob','f4a644cc55909f01a14fcffe5b7b430048c758ad')]
for row in notes:n.append(row)
n.column_dimensions['A'].width=20;n.column_dimensions['B'].width=110
for row in n:
 for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
for sheet in wb:
 for c in sheet[1]:c.font=Font(name='Arial',size=12,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='244F73')
 for row in sheet.iter_rows(min_row=2):
  for c in row:
   c.font=Font(name='Arial' if sheet!=n else 'Microsoft YaHei',size=11)
   if isinstance(c.value,float):c.number_format='0.000000' if sheet in [s,t] else '0.00000000'
 for i in range(1,sheet.max_row+1):sheet.row_dimensions[i].height=28 if sheet!=n else 42
path=d/'VEGFA_TPM_plot_data.xlsx';wb.save(path)
with (d/'VEGFA_TPM_Prism.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['Normoxia','Hypoxia'])
 for a,b in zip(nor,hyp):w.writerow([a['TPM'],b['TPM']])
v=load_workbook(path,data_only=True)
assert list(v['Prism_columns'].values)==[('Normoxia','Hypoxia')]+[(float(a['TPM']),float(b['TPM'])) for a,b in zip(nor,hyp)]
print(json.dumps({'xlsx':str(path),'means':{k:str(v) for k,v in means.items()},'ratio':str(means['Hypoxia']/means['Normoxia']),'samples':len(rows)}))
