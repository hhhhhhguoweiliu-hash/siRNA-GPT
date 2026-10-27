"""Reproduce charts from two untouched submitted rating forms; no adjudication.
Usage: python make_expert_figures.py /absolute/path/results_original.json
"""
from pathlib import Path
from collections import Counter
from functools import lru_cache
from xml.sax.saxutils import escape
import sys,json,hashlib,re,csv,math,shutil
from openpyxl import load_workbook,Workbook
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image,ImageDraw,ImageFont
from pypdf import PdfReader,PdfWriter

ROOT=Path(__file__).resolve().parent;FIG=ROOT/'figures';FIG.mkdir(exist_ok=True)
FONT=Path(r'C:\Windows\Fonts\arialbd.ttf');assert FONT.exists()
pdfmetrics.registerFont(TTFont('ArialBold',str(FONT)))
RAW=Path(sys.argv[1]);results=json.loads(RAW.read_text(encoding='utf-8'));KEY={r['answer_id']:r for r in results}
assert hashlib.sha256(RAW.read_bytes()).hexdigest()=='f443d6f358a3c63f45eeef0906ae968f74f231534eb3c5276207767a14c2b9a3'
GROUPS=['llm','workflow','agent'];NAMES={'llm':'LLM','workflow':'Fixed workflow','agent':'Agent'}
COLORS={'llm':'#8295B3','workflow':'#E6B77C','agent':'#A44958'}
DIMS=['Factual\naccuracy','Evidence\nsupport','Requirement\ncompletion','Boundary\nhandling']
KEYDIMS=['factual_accuracy','evidence_support','requirement_completion','boundary_handling']
STRATA={'Standard tasks':['W1','W2','W2b','W3'],'Requirement change':['U1'],'Reliability':['R1','R2','R3','R4']}
A=ROOT/'source_form_A_original.xlsx'
B=ROOT/'source_completion_scores_original.xlsx'
assert hashlib.sha256(A.read_bytes()).hexdigest()=='1463b5641298bee179b7c05b9f8b08697afdf856a055c57af7411ea9c995b675'
assert hashlib.sha256(B.read_bytes()).hexdigest()=='97759996c7a8a2805e914fac6ee09d85428e35ff7e22a00de2bfa5318c84bcf4'
rows=[]
wa=load_workbook(A,data_only=False)
for rn,v in enumerate(wa['评分总表'].iter_rows(values_only=True),1):
 if len(v)>9 and v[3] in KEY:
  aid=v[3];rows.append(dict(form='A',answer_id=aid,group=KEY[aid]['group'],case_id=KEY[aid]['case_id'],scores=list(v[4:8]),completion_text=v[9],source_sheet='评分总表',source_row=rn,reason=None))
wb=load_workbook(B,data_only=True)
for rn,v in enumerate(wb.active.iter_rows(values_only=True),1):
 m=re.search(r'\bA\d{3}\b',str(v[1] if len(v)>1 else ''))
 if not m:continue
 aid=m.group(0);assert aid in KEY
 assert v[6]==sum(v[2:6]),aid
 rows.append(dict(form='B',answer_id=aid,group=KEY[aid]['group'],case_id=KEY[aid]['case_id'],scores=list(v[2:6]),completion_text=v[7],source_sheet=wb.active.title,source_row=rn,reason=v[8]))
for f in ['A','B']:
 rr=[r for r in rows if r['form']==f]
 assert len(rr)==27 and {r['answer_id'] for r in rr}==set(KEY)
for r in rows:
 assert all(type(v) is int and 0<=v<=2 for v in r['scores'])
 t=r['completion_text']
 r['completion']='partial' if t.startswith('部分完成') else ('not_completed' if t.startswith(('未交付','未完成')) else ('completed' if t.startswith('完成') else None))
 assert r['completion'] is not None
 r['stratum']=next(k for k,v in STRATA.items() if r['case_id'] in v)
means=[];completion=[]
for f in ['A','B']:
 for st,cases in STRATA.items():
  for g in GROUPS:
   rr=[r for r in rows if r['form']==f and r['stratum']==st and r['group']==g];assert len(rr)==len(cases)
   for i,d in enumerate(KEYDIMS):
    means.append(dict(form=f,stratum=st,group=g,dimension=d,n=len(rr),mean=sum(r['scores'][i] for r in rr)/len(rr)))
   count=Counter(r['completion'] for r in rr);n=len(rr);k=count['completed']
   completion.append(dict(form=f,stratum=st,group=g,n=n,completed=k,partial=count['partial'],not_completed=count['not_completed'],percent=100*k/n))
def savecsv(name,rs):
 with (ROOT/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rs[0]));w.writeheader();w.writerows(rs)
long=[dict(form=r['form'],answer_id=r['answer_id'],case_id=r['case_id'],group=r['group'],stratum=r['stratum'],dimension=d,score=r['scores'][i],source_sheet=r['source_sheet'],source_row=r['source_row']) for r in rows for i,d in enumerate(KEYDIMS)]
savecsv('scores_long.csv',long);savecsv('dimension_means.csv',means);savecsv('completion_as_submitted.csv',completion)
diff=[]
for f in ['A','B']:
 for st in STRATA:
  for a,b in [('agent','llm'),('workflow','llm'),('agent','workflow')]:
   x=next(r for r in completion if r['form']==f and r['stratum']==st and r['group']==a)
   y=next(r for r in completion if r['form']==f and r['stratum']==st and r['group']==b)
   diff.append(dict(form=f,stratum=st,comparison=a+'_minus_'+b,n_each=x['n'],percentage_point_difference=x['percent']-y['percent'],status='AS_SUBMITTED_NOT_ADJUDICATED'))
savecsv('completion_differences_pp.csv',diff)
wout=Workbook();wout.remove(wout.active)
for name,records in [('Scores_long',long),('Dimension_means',means),('Completion',completion),('Completion_diff_pp',diff)]:
 ws=wout.create_sheet(name);ws.append(list(records[0]))
 for r in records:ws.append(list(r.values()))
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for col in ws.columns:
  letter=col[0].column_letter;ws.column_dimensions[letter].width=min(45,max(16,max(len(str(c.value or '')) for c in col[:50])+2))
meta=wout.create_sheet('Read_me')
for row in [
 ['status','Submitted ratings; source discrepancies unresolved; forms not pooled.'],
 ['form_A','siRNA_GPT_三组问答_专家评分表_最终版.xlsx'],
 ['form_B','专家评分-任务完成情况评分.xlsx'],
 ['raters','Two forms are not assumed to be two independent raters.'],
 ['unit','9 seen development templates x 3 groups, not 27 independent questions.'],
 ['plots','Means and all task dots. No trial/rater error bars or significance tests.'],
 ['completion','Original labels only. Completed does not mean biological efficacy.'],
 ['author_changes','0 expert scores changed; 0 new model calls.']]:meta.append(row)
meta.column_dimensions['A'].width=24;meta.column_dimensions['B'].width=105
wout.save(ROOT/'expert_plot_data.xlsx')

class Plot:
 def __init__(self,w,h):self.w=w;self.h=h;self.items=[]
 def line(self,x1,y1,x2,y2,color='#262626',width=1):self.items.append(('line',x1,y1,x2,y2,color,width))
 def rect(self,x,y,w,h,color,stroke=None):self.items.append(('rect',x,y,w,h,color,stroke))
 def dot(self,x,y,r=3,color='#262626',fill='#FFFFFF'):self.items.append(('dot',x,y,r,color,fill))
 def text(self,x,y,t,size=17,anchor='start',color='#262626',angle=0):self.items.append(('text',x,y,str(t),size,anchor,color,angle))
 def output(self,name,width_mm=180):
  s=width_mm*72/25.4/self.w;pdf=canvas.Canvas(str(FIG/(name+'.pdf')),pagesize=(self.w*s,self.h*s));svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{width_mm}mm" height="{width_mm*self.h/self.w}mm" viewBox="0 0 {self.w} {self.h}">','<rect width="100%" height="100%" fill="white"/>']
  scale=width_mm/25.4*600/self.w;im=Image.new('RGB',(round(self.w*scale),round(self.h*scale)),'white');dr=ImageDraw.Draw(im)
  @lru_cache(None)
  def ft(sz):return ImageFont.truetype(str(FONT),round(sz*scale))
  for q in self.items:
   if q[0]=='line':
    _,x1,y1,x2,y2,c,lw=q;pdf.setStrokeColor(c);pdf.setLineWidth(lw*s);pdf.line(x1*s,(self.h-y1)*s,x2*s,(self.h-y2)*s);svg.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="{c}" stroke-width="{lw}"/>');dr.line((x1*scale,y1*scale,x2*scale,y2*scale),fill=c,width=max(1,round(lw*scale)))
   elif q[0]=='rect':
    _,x,y,w,h,c,stroke=q
    if h<0:raise ValueError('negative bar')
    pdf.setFillColor(c);pdf.setStrokeColor(stroke or c);pdf.setLineWidth(.65*s);pdf.rect(x*s,(self.h-y-h)*s,w*s,h*s,fill=1,stroke=bool(stroke));svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{c}" stroke="{stroke or c}" stroke-width=".65"/>');dr.rectangle((x*scale,y*scale,(x+w)*scale,(y+h)*scale),fill=c,outline=stroke,width=max(1,round(.65*scale)))
   elif q[0]=='dot':
    _,x,y,r,c,fill=q;pdf.setFillColor(fill);pdf.setStrokeColor(c);pdf.setLineWidth(.9*s);pdf.circle(x*s,(self.h-y)*s,r*s,fill=1,stroke=1);svg.append(f'<circle cx="{x}" cy="{y}" r="{r}" fill="{fill}" stroke="{c}" stroke-width=".9"/>');dr.ellipse(((x-r)*scale,(y-r)*scale,(x+r)*scale,(y+r)*scale),fill=fill,outline=c,width=max(1,round(.9*scale)))
   else:
    _,x,y,t,sz,anchor,c,angle=q
    for j,txt in enumerate(t.split('\n')):
     yy=y+j*sz*1.17;off={'start':0,'middle':-.5,'end':-1}[anchor]*pdfmetrics.stringWidth(txt,'ArialBold',sz)
     pdf.saveState();pdf.translate(x*s,(self.h-yy)*s);pdf.rotate(-angle);pdf.setFillColor(c);pdf.setFont('ArialBold',sz*s);pdf.drawString(off*s,0,txt);pdf.restoreState()
     rot=f' transform="rotate({angle} {x} {yy})"' if angle else ''
     svg.append(f'<text x="{x}" y="{yy}" font-family="Arial" font-weight="700" font-size="{sz}" text-anchor="{anchor}" fill="{c}"{rot}>{escape(txt)}</text>')
     if not angle:dr.text((x*scale,yy*scale),txt,font=ft(sz),fill=c,anchor={'start':'ls','middle':'ms','end':'rs'}[anchor])
     else:
      box=ft(sz).getbbox(txt);patch=Image.new('RGBA',(box[2]-box[0]+20,box[3]-box[1]+20),(255,255,255,0));pd=ImageDraw.Draw(patch);pd.text((10-box[0],10-box[1]),txt,font=ft(sz),fill=c);patch=patch.rotate(-angle,expand=True,resample=Image.Resampling.BICUBIC);im.paste(patch,(round(x*scale-patch.width/2),round(yy*scale-patch.height/2)),patch)
  pdf.showPage();pdf.save();svg.append('</svg>');(FIG/(name+'.svg')).write_text('\n'.join(svg),encoding='utf-8');im.save(FIG/(name+'.png'),dpi=(600,600))

def legend(p,y=74):
 for x,g in zip([300,505,795],GROUPS):
  p.rect(x,y-14,20,14,COLORS[g]);p.text(x+29,y,NAMES[g],17)
def axes(p,x,y,w,h,ymax,ticks,label):
 for v in ticks:
  yy=y+h-h*v/ymax;p.line(x,yy,x+w,yy,'#E7E7E7',.75);p.line(x-5,yy,x,yy);p.text(x-12,yy+5,str(v),15,'end')
 p.line(x,y,x,y+h);p.line(x,y+h,x+w,y+h);p.text(x-49,y+h/2,label,17,'middle',angle=-90)
def quality(st,name):
 p=Plot(1100,520);p.text(42,36,'Expert assessment: '+st.lower(),26);p.text(42,58,'Two submitted forms shown separately',15,color='#666666');legend(p,91)
 for pi,f in enumerate(['A','B']):
  x=86+pi*550;y=155;ww=440;hh=225
  p.text(x-48,131,chr(97+pi),23);p.text(x+ww/2,131,'Form '+f,21,'middle')
  axes(p,x,y,ww,hh,2,[0,1,2],'Expert score (0–2)')
  for di,d in enumerate(KEYDIMS):
   center=x+(di+.5)*ww/4
   for gi,g in enumerate(GROUPS):
    value=next(r['mean'] for r in means if r['form']==f and r['stratum']==st and r['group']==g and r['dimension']==d)
    xx=center+(gi-1)*27;bh=hh*value/2;p.rect(xx-10,y+hh-bh,20,bh,COLORS[g],'#414141')
    vals=[r['scores'][di] for r in rows if r['form']==f and r['stratum']==st and r['group']==g]
    for vi,v in enumerate(vals):p.dot(xx+(vi-(len(vals)-1)/2)*3.7,y+hh-hh*v/2,2.8)
   p.text(center,y+hh+29,DIMS[di],16,'middle')
 p.text(42,467,f'Dots: all task scores; bars: means. n = {len(STRATA[st])} task episode(s) per group.',16)
 p.text(42,494,'Seen development cases. Rating–evidence discrepancies remain unresolved; forms are not pooled.',14,color='#777777')
 p.output(name)
quality('Standard tasks','fig1_standard_quality')
quality('Reliability','figS1_reliability_quality')
quality('Requirement change','figS2_requirement_change_quality')
p=Plot(1200,780);p.text(44,35,'Task completion in submitted expert forms',27);p.text(44,60,'Original completion labels; each bar shows completed task episodes / available episodes',16);legend(p,91)
for fi,f in enumerate(['A','B']):
 for si,(st,cases) in enumerate(STRATA.items()):
  x=90+si*400;y=175+fi*300;ww=295;hh=170
  p.text(x-48,y-30,chr(97+fi*3+si),22);p.text(x+ww/2,y-30,st+f' (n = {len(cases)})',18,'middle')
  axes(p,x,y,ww,hh,100,[0,50,100],'Completed tasks (%)')
  for gi,g in enumerate(GROUPS):
   rr=next(r for r in completion if r['form']==f and r['stratum']==st and r['group']==g)
   xx=x+(gi+.5)*ww/3;bh=hh*rr['percent']/100;p.rect(xx-24,y+hh-bh,48,bh,COLORS[g],'#414141')
   p.text(xx,y+hh-bh-10,f"{rr['completed']}/{rr['n']}",18,'middle')
   p.text(xx,y+hh+29,{'llm':'LLM','workflow':'Fixed\nworkflow','agent':'Agent'}[g],17,'middle')
 p.text(44,fi*300+116,'Form '+f,16,color='#666666')
p.text(44,734,'Seen development cases; n = 1 for requirement change. No trial error bars or significance claims.',16)
p.text(44,762,'As submitted; source discrepancies unresolved. These are not adjudicated accuracy or efficacy rates.',14,color='#777777')
p.output('fig2_completion_as_submitted')
writer=PdfWriter()
for name in ['fig1_standard_quality','fig2_completion_as_submitted','figS1_reliability_quality','figS2_requirement_change_quality']:
 for pg in PdfReader(FIG/(name+'.pdf')).pages:writer.add_page(pg)
with (ROOT/'siRNA_GPT_expert_figures_REVIEW_PENDING.pdf').open('wb') as f:writer.write(f)
checks={'source_A_sha256':hashlib.sha256(A.read_bytes()).hexdigest(),'source_B_sha256':hashlib.sha256(B.read_bytes()).hexdigest(),'forms':2,'independent_raters_assumed':False,'scores_preserved':216,'full_answer_positions_per_form':27,'new_scores_or_adjudications':0,'font_file':str(FONT),'all_plot_text_font':'Arial Bold','png_dpi':600,'pdf_font_embedding':[]}
for pth in sorted(FIG.glob('*.pdf')):
 rd=PdfReader(pth);fonts=[]
 for page in rd.pages:
  for ob in page['/Resources']['/Font'].get_object().values():
   fo=ob.get_object();fd=fo.get('/FontDescriptor');fonts.append({'name':str(fo.get('/BaseFont')),'embedded':bool(fd and '/FontFile2' in fd.get_object())})
 bold=[x for x in fonts if 'Arial' in x['name']];assert bold and all(x['embedded'] and 'Bold' in x['name'] for x in bold)
 checks['pdf_font_embedding'].append({'file':pth.name,'fonts':fonts})
checks['complete_by_form_group']={f:{g:sum(r['completed'] for r in completion if r['form']==f and r['group']==g) for g in GROUPS} for f in ['A','B']}
(ROOT/'figure_checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(ROOT/'plot_data.json').write_text(json.dumps({'rows':rows,'means':means,'completion':completion,'source_status':'USER_SUBMITTED_EXPERT_FORMS_UNADJUDICATED','form_relationship':'not clarified; never pooled'},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'figure_count':4,'pages':len(PdfReader(ROOT/'siRNA_GPT_expert_figures_REVIEW_PENDING.pdf').pages),'completion':completion,'font_checks_passed':True},ensure_ascii=True))

