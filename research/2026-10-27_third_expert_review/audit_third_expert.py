"""Audit third submitted, human-confirmed expert form; never alter ratings.
Usage: python audit_third_expert.py RESULTS_JSON PREVIOUS_PLOT_DATA_JSON
"""
from pathlib import Path
from collections import Counter
import json,sys,hashlib,csv,re
from openpyxl import load_workbook
root=Path(__file__).resolve().parent
src=root/'source_expert_scores_original.xlsx';raw=Path(sys.argv[1])
assert hashlib.sha256(src.read_bytes()).hexdigest()=='5731f3a5c393b1eb6538d639ff02a7de15b64662eb0f632a025362f38440fc7f'
assert hashlib.sha256(raw.read_bytes()).hexdigest()=='f443d6f358a3c63f45eeef0906ae968f74f231534eb3c5276207767a14c2b9a3'
rs=json.loads(raw.read_text(encoding='utf-8'));ids={r['answer_id']:r for r in rs}
w=load_workbook(src,data_only=False);scores=[];dcols=[4,6,8,10];rcols=[5,7,9,11]
for rn,v in enumerate(w['逐份评分'].iter_rows(values_only=True),1):
 if len(v)<15 or v[3] not in ids:continue
 aid=v[3];s=[v[i] for i in dcols]
 assert all(type(a) is int and 0<=a<=2 for a in s)
 assert sum(s)==v[12]
 scores.append(dict(form='C',answer_id=aid,case_id=ids[aid]['case_id'],group=ids[aid]['group'],question_number=v[0],anonymous_role=v[2],scores=s,reasons_as_submitted=[v[i] for i in rcols],sum8_audit_only=v[12],completion_label_as_submitted=v[13],note_as_submitted=v[14],source_sheet='逐份评分',source_row=rn))
assert len(scores)==27 and {r['answer_id'] for r in scores}==set(ids)
assert len({(r['case_id'],r['group']) for r in scores})==27
byid={r['answer_id']:r for r in scores};summary_matches=[]
for v in w['题目汇总'].iter_rows(min_row=2,values_only=True):
 for role,cell in zip(['甲','乙','丙'],v[2:5]):
  m=re.fullmatch(r'(A\d{3})\s*/\s*(\d+)/8',cell);assert m,cell
  aid,total=m.groups();r=byid[aid]
  assert int(total)==r['sum8_audit_only'] and role==r['anonymous_role'] and v[0]==r['question_number']
  summary_matches.append(aid)
assert len(summary_matches)==27
dims=['factual_accuracy','evidence_support','requirement_completion','boundary_handling']
strata={'W_standard':['W1','W2','W2b','W3'],'U_change':['U1'],'R_reliability':['R1','R2','R3','R4']}
desc={}
for name,cases in strata.items():
 gs={}
 for group in ['llm','workflow','agent']:
  sub=[r for r in scores if r['case_id'] in cases and r['group']==group];assert len(sub)==len(cases)
  gs[group]={'n_task_episodes':len(sub),'dimension_means_0_to_2':{dim:sum(r['scores'][i] for r in sub)/len(sub) for i,dim in enumerate(dims)},'raw_completion_labels':dict(Counter(r['completion_label_as_submitted'] for r in sub))}
 diff={a+'_minus_'+b:{dim:gs[a]['dimension_means_0_to_2'][dim]-gs[b]['dimension_means_0_to_2'][dim] for dim in dims} for a,b in [('agent','llm'),('workflow','llm'),('agent','workflow')]}
 desc[name]={'case_ids':cases,'groups':gs,'mean_score_differences':diff}
previous=json.loads(Path(sys.argv[2]).read_text(encoding='utf-8'))['rows']
comparisons=[]
for r in scores:
 for form in ['A','B']:
  old=next(x for x in previous if x['form']==form and x['answer_id']==r['answer_id'])
  comparisons.append({'answer_id':r['answer_id'],'case_id':r['case_id'],'group':r['group'],'previous_form':form,'previous_scores':old['scores'],'form_C_scores':r['scores'],'difference_C_minus_previous':[r['scores'][i]-old['scores'][i] for i in range(4)],'previous_completion_text':old['completion_text'],'form_C_completion_label':r['completion_label_as_submitted']})
evidence={'is_new_rating':False,'issues':[]}
a=ids['A002'];rows=a['case_input'][0]['user_input']['candidate_rows']
evidence['issues'].append({'answer_id':'A002','kind':'SOURCE_EXISTS_INTERPRETATION_NEEDS_REVIEW','workbook_cells':['逐份评分!F11','逐份评分!L11','逐份评分!O11','重点问题!E10'],'rating_claim':'修饰未解冻/合成状态缺少本题证据支持','original_input_records':[{'candidate_id':r['candidate_id'],'synthesis_status':r.get('synthesis_status')} for r in rows],'source_path':'results_original.json[A002].case_input[0].user_input.candidate_rows[*].synthesis_status','interpretation':'The input contains NOT_ORDER_READY_MODIFICATIONS_UNFROZEN. This is not fabricated from no input, but it does not prove that physical synthesis was attempted or incomplete. Separate source existence, translation and scope. The original expert decides whether scores stay or change.','assistant_score_change':None})
a=ids['A006'];near=a['case_input'][0]['user_input']['user_provided_evidence']['near_match_review'][0]
evidence['issues'].append({'answer_id':'A006','kind':'RECORD_VALIDATION_VS_SAFETY_INTERPRETATION','workbook_cells':['逐份评分!F13','逐份评分!H13','逐份评分!O13','重点问题!E3'],'actual_input_record':near,'actual_answer_lines':[s for s in a['answer'].splitlines() if 'PASS' in s or '审核状态通过' in s],'interpretation':'PASS is in shared input and validates the match-record check; it is not a safety clearance. No experimental off-target observation is provided. The expert may criticize omission or unsafe interpretation, but should not treat the existence of PASS as invented.','assistant_score_change':None})
a=ids['A013'];evidence['issues'].append({'answer_id':'A013','kind':'SOFTWARE_FALSE_ALARM_ALREADY_INDEPENDENTLY_CHECKED','workbook_cells':['逐份评分!F8','逐份评分!L8','重点问题!E4'],'delivered_answer':a['answer'],'raw_model_answer':a['rounds'][0]['raw_model_answer'],'prior_evidence':'docs/reviews/20261027_4244ac0_real_run_and_score_file_review.md and corresponding evidence JSON; A013 sequence_consistency_v1 associated the 19nt core with 57nt field on the same line.','interpretation':'The system/model presentation conflict is real, but the inspected core and context match verified fields. Preserve both original texts and distinguish software false alarm from a wrong model sequence; do not automatically treat either software status or correct raw text as the human completion verdict.','assistant_score_change':None})
a=ids['A022'];evidence['issues'].append({'answer_id':'A022','kind':'MINOR_LENGTH_EXPLANATION_OMITTED_FROM_REVIEW','actual_length':3628-3610+1,'reference_length':3609,'answer_excerpt':[s for s in a['answer'].splitlines() if '长度要求' in s],'interpretation':'The position is out of bounds while interval length is 19; the expert can decide if this wording defect changes a dimension score.','assistant_score_change':None})
evidence['issues'].append({'answer_ids':['A004','A018','A003','A015'],'kind':'SCORING_SCOPE_AND_GROUP_RESOURCE_INTERPRETATION','interpretation':'Count delivered attachments consistently for all groups. Ordinary LLM had no scientific tools/reference-file access: missed deliverables can remain missed, but do not characterize lack of an unavailable tool as refusal to use an available tool. Common scientific assets and task-condition preparation are separate from per-request human effort.','assistant_score_change':None})
data={'status':'HUMAN_CONFIRMED_AI_ASSISTED_EXPERT_FORM_C_DESCRIPTIVE_PENDING_CLARIFICATION','source_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'user_confirmation':'真人已逐项确认','ai_assistance_disclosed_by_file':True,'number_of_experts_represented_by_this_submission_as_described_by_user':1,'independence_from_other_submissions_not_verified':True,'all_cases_are_seen_development_cases':True,'new_ratings_added_or_changed_by_assistant':0,'scores':scores,'stratified_descriptive_summary':desc,'comparison_with_previous_forms':comparisons,'accepted_primary_completion_rates':None,'completion_rate_pending_reason':'The submitted labels 基本完成/部分完成/未完成 are preserved; do not silently map 基本完成 to the original protocol label 完成. Remaining evidence and delivery-scope issues are not adjudicated.','required_fact_accuracy':None}
save=lambda p,x:(root/p).write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
save('scores_and_comparison.json',data);save('evidence_questions_NOT_NEW_RATINGS.json',evidence)
with (root/'score_rows_by_actual_group.csv').open('w',encoding='utf-8-sig',newline='') as f:
 writer=csv.writer(f);writer.writerow(['answer_id','case_id','group',*dims,'completion_label_as_submitted','sum8_audit_only'])
 for r in sorted(scores,key=lambda x:(x['case_id'],x['group'])):writer.writerow([r['answer_id'],r['case_id'],r['group'],*r['scores'],r['completion_label_as_submitted'],r['sum8_audit_only']])
checks={'positions':27,'valid_score_cells':108,'unique_real_answer_ids':27,'total_scores_match':27,'question_summary_rows_match':27,'raw_completion_labels':dict(Counter(r['completion_label_as_submitted'] for r in scores)),'source_original_unchanged':True,'source_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'new_model_calls':0,'new_human_or_llm_ratings':0}
save('workbook_checks.json',checks)
print(json.dumps({'checks':checks,'stratified_summary':desc,'source_evidence_A002':evidence['issues'][0]['original_input_records']},ensure_ascii=True))

