"""Read one user-confirmed expert workbook; preserve every original rating."""
import argparse
import collections
import csv
import hashlib
import json
import re
import shutil
from copy import copy
from pathlib import Path

import openpyxl

DIMENSIONS = ['factual_accuracy', 'evidence_support', 'requirement_completion', 'boundary_handling']
GROUPS = ['llm', 'workflow', 'agent']
CATEGORY = {'W1':'standard','W2':'standard','W2b':'standard','W3':'standard',
            'U1':'requirement_change','R1':'reliability','R2':'reliability','R3':'reliability','R4':'reliability'}
EXPECTED_SHA = '2c0658ca713634fc25489ab1123ebdbc159f107d04fc2f8698b1a78685d7d45b'


def average(values):
    return sum(values) / len(values)


def summarize(rows, category):
    return [dict(category=category,group=g,n_tasks=len(r),
                 sum_total=sum(x['total'] for x in r),
                 mean_total=average([x['total'] for x in r]),
                 **{k:average([x[k] for x in r]) for k in DIMENSIONS})
            for g in GROUPS if (r := [x for x in rows if x['group']==g])]


def write_csv(path, rows):
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main():
    p=argparse.ArgumentParser()
    p.add_argument('--source',type=Path,required=True)
    p.add_argument('--original-source',type=Path)
    p.add_argument('--mapping',type=Path,required=True)
    p.add_argument('--origin-audit',type=Path,required=True)
    p.add_argument('--output',type=Path,required=True)
    args=p.parse_args()
    source_bytes=args.source.read_bytes()
    original_path=args.original_source or args.source
    original=original_path.read_bytes()
    digest=hashlib.sha256(source_bytes).hexdigest()
    assert hashlib.sha256(original).hexdigest()==EXPECTED_SHA,'Unexpected original file; do not mix rater versions'
    correction=args.original_source is not None
    mapping=json.loads(args.mapping.read_text(encoding='utf-8-sig'))
    key={x['answer_id']:x for x in mapping}
    assert len(key)==27
    wb=openpyxl.load_workbook(args.source,data_only=False)
    cached=openpyxl.load_workbook(args.source,data_only=True)
    formatting_changes=[]
    if correction:
        ow=openpyxl.load_workbook(original_path,data_only=False)
        assert wb.sheetnames==ow.sheetnames
        value_changes=[]
        for row in ow.active:
            for cell in row:
                edited=wb.active[cell.coordinate]
                if cell.value!=edited.value:
                    value_changes.append((cell.coordinate,cell.value,edited.value))
                if any(copy(getattr(cell,k))!=copy(getattr(edited,k)) for k in ['font','fill','alignment','border','protection']) or cell.number_format!=edited.number_format:
                    formatting_changes.append(cell.coordinate)
        assert value_changes==[('D27',1,2),('E27',1,2)],value_changes
        assert not formatting_changes,formatting_changes
    assert wb.sheetnames==['Sheet1']
    ws=wb['Sheet1'];cw=cached['Sheet1']
    assert (ws.max_row,ws.max_column)==(28,7)
    assert [ws.cell(1,j).value for j in range(1,8)]==['题号','答复（登记号）','事实准确','证据支持','需求完成','边界处理','合计/8']
    rows=[];cache_matches=0
    for i in range(2,29):
        q,ident,*_=list(ws.iter_rows(min_row=i,max_row=i,values_only=True))[0]
        match=re.fullmatch(r'([甲乙丙])\s+(A\d{3})',ident)
        assert match,(i,ident)
        letter,aid=match.groups();info=key[aid]
        assert q==info['question_number'] and letter==info['display_answer']
        scores=[ws.cell(i,j).value for j in range(3,7)]
        assert all(type(v) is int and 0<=v<=2 for v in scores),(aid,scores)
        formula='='+ '+'.join(f'{col}{i}' for col in 'CDEF')
        assert ws.cell(i,7).value==formula
        total=sum(scores)
        assert cw.cell(i,7).value==total,(aid,cw.cell(i,7).value,total)
        cache_matches+=1
        rows.append(dict(question_number=q,case_id=info['case_id'],question_name=info['question_name'],
                         display_answer=letter,answer_id=aid,group=info['group'],category=CATEGORY[info['case_id']],
                         **dict(zip(DIMENSIONS,scores)),total=total,source_cells=f'Sheet1!C{i}:F{i}',
                         source_total_formula_cell=f'Sheet1!G{i}'))
    assert len({x['answer_id'] for x in rows})==27
    assert set(x['answer_id'] for x in rows)==set(key)
    assert collections.Counter(x['group'] for x in rows)==dict.fromkeys(GROUPS,9)
    assert collections.Counter(x['question_number'] for x in rows)==dict.fromkeys(range(1,10),3)
    summary=summarize(rows,'all')
    for category in ['standard','requirement_change','reliability']:
        summary+=summarize([x for x in rows if x['category']==category],category)
    bycase={q:{x['group']:x for x in rows if x['question_number']==q} for q in range(1,10)}
    paired=[]
    for q,r in bycase.items():
        paired.append(dict(question_number=q,case_id=r['agent']['case_id'],
                           llm=r['llm']['total'],workflow=r['workflow']['total'],agent=r['agent']['total'],
                           agent_minus_llm=r['agent']['total']-r['llm']['total'],
                           agent_minus_workflow=r['agent']['total']-r['workflow']['total']))
    comparisons={}
    for baseline in ['llm','workflow']:
        diffs=[x['agent_minus_'+baseline] for x in paired]
        avg_base=next(x['mean_total'] for x in summary if x['category']=='all' and x['group']==baseline)
        comparisons[baseline]={'paired_wins':sum(x>0 for x in diffs),'paired_ties':sum(x==0 for x in diffs),
            'paired_losses':sum(x<0 for x in diffs),'mean_total_difference':average(diffs),
            'relative_total_rating_difference':average(diffs)/avg_base,
            'meaning':'Equal-weight 0-8 ordinal rating difference, not classification accuracy or a biological effect'}
    origin=json.loads(args.origin_audit.read_text(encoding='utf-8-sig'))
    assert origin['R4']['answer_files_byte_identical'] is True
    assert origin['A013']['body_contains_scaled_value'] is False
    byid={x['answer_id']:x for x in rows}
    q9_difference={k:byid['A025'][k]-byid['A011'][k] for k in DIMENSIONS}
    result={'review_date':'2026-10-28','timezone':'Asia/Shanghai',
        'reviewed_commit':'888804378d5f62a1d0093f9da230a466ed67daf8',
        'source_filename':args.source.name,'source_sha256':digest,'source_bytes':len(source_bytes),
        'original_upload_sha256':EXPECTED_SHA,'original_upload_preserved':True,
        'user_transcription_correction_applied':correction,
        'reviewer_provenance':'User explicitly confirms this file was scored by one authoritative human expert',
        'rater_count':1,'blind_review_method_recorded_in_workbook':False,
        'new_assistant_quality_scores':0,'new_model_calls':0,
        'n_answers':27,'n_cases_per_group':9,'summary':summary,'paired_totals':paired,
        'comparisons':comparisons,
        'R1_R2_share_of_agent_workflow_total_point_gap':sum(x['agent_minus_workflow'] for x in paired if x['case_id'] in ['R1','R2'])/sum(x['agent_minus_workflow'] for x in paired),
        'resolved_recording_issue':{'case_id':'R4','answer_ids':['A011','A025'],
            'current_totals':[byid['A011']['total'],byid['A025']['total']],
            'agent_minus_workflow_by_dimension':q9_difference,
            'resolution':'User confirmed both are 8; original difference was a transcription error',
            'correction_applied':correction},
        'interpretation_items_not_score_changes':[
            {'case_id':'W2b','answer_id':'A013','original_total':byid['A013']['total'],
             'needed':'State whether raw model text/attachments were counted; do not infer evaluator intent or replace delivered body'}],
        'classification_metrics':{'accuracy':None,'precision':None,'recall':None,'f1':None},
        'formal_task_success_rate':None,'p_values':None,
        'limitations':['One human rater','Nine seen, partly related development templates per configuration',
            'No per-fact labels, pass/fail task judgements or score rationales in this workbook',
            'Template, shared validation and model contributions remain distinct from aggregate system ratings']}
    checks={'sheet_count':len(wb.sheetnames),'score_rows':len(rows),'valid_original_dimension_scores':108,
        'cached_totals_match_independent_sums':cache_matches,'mapping_matches':27,
        'comments':sum(c.comment is not None for sh in wb for row in sh for c in row),
        'original_sha256':EXPECTED_SHA,'original_unchanged':original_path.read_bytes()==original,
        'current_sha256':digest,'user_correction_applied':correction,
        'changed_score_cells':['Sheet1!D27','Sheet1!E27'] if correction else [],
        'formatting_changes':formatting_changes,
        'quality_rescoring_by_assistant':False}
    args.output.mkdir(parents=True,exist_ok=True)
    copied=args.output/'source_workbook_original.xlsx'
    if copied.resolve()!=original_path.resolve():shutil.copyfile(original_path,copied)
    assert copied.read_bytes()==original
    if correction:
        revised=args.output/'source_workbook_user_corrected.xlsx'
        if revised.resolve()!=args.source.resolve():shutil.copyfile(args.source,revised)
        assert revised.read_bytes()==source_bytes
    write_csv(args.output/'scores_by_actual_group.csv',rows)
    write_csv(args.output/'group_summary.csv',summary)
    for name,data in [('analysis.json',result),('workbook_checks.json',checks)]:
        (args.output/name).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'summary':summary,'comparisons':comparisons,'paired':paired,'checks':checks},ensure_ascii=False))


if __name__=='__main__':main()
