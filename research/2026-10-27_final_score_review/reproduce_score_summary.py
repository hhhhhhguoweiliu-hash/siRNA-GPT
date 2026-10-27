"""Extract submitted ratings verbatim; calculate descriptive means only.
Usage: python reproduce_score_summary.py RESULTS_JSON [PREVIOUS_XLSX]
Requires openpyxl. No model calls, new ratings, or source modifications.
"""
from pathlib import Path
from collections import Counter
import json,hashlib,csv,sys
from openpyxl import load_workbook
root=Path(__file__).resolve().parent
src=root/'source_workbook_original.xlsx'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(src)=='1463b5641298bee179b7c05b9f8b08697afdf856a055c57af7411ea9c995b675'
raw=Path(sys.argv[1])
assert sha(raw)=='f443d6f358a3c63f45eeef0906ae968f74f231534eb3c5276207767a14c2b9a3'
results=json.loads(raw.read_text(encoding='utf-8'))
key={r['answer_id']:r for r in results}
dims=['factual_accuracy','evidence_support','requirement_completion','boundary_handling']
zh=['事实准确','证据支持','需求完成','边界处理']
strata={'W_standard':['W1','W2','W2b','W3'],'U_change':['U1'],'R_reliability':['R1','R2','R3','R4']}
def read(w):
 rows=[]
 for cells in w['评分总表']:
  v=[c.value for c in cells]
  if len(v)<10 or v[3] not in key:continue
  assert all(type(x) is int and 0<=x<=2 for x in v[4:8])
  aid=v[3]
  rows.append(dict(answer_id=aid,case_id=key[aid]['case_id'],group=key[aid]['group'],question_number=v[0],anonymous_role=v[2],scores=v[4:8],completion_label_as_submitted=v[9],sum8_audit_only=sum(v[4:8]),score_row=cells[0].row,total_formula=v[8]))
 return rows
wb=load_workbook(src,data_only=False);rows=read(wb);ids={r['answer_id']:r for r in rows}
assert len(rows)==len(ids)==len(key)==27 and set(ids)==set(key)
assert len({(r['group'],r['case_id']) for r in rows})==27
for r in rows:
 n=r['score_row'];assert r['total_formula']==f'=SUM(E{n}:H{n})'
reasons=list(wb['评分理由'].iter_rows(min_row=2,values_only=True));assert len(reasons)==27
for aid,question,role,scoretext,total,reason in reasons:
 assert scoretext=='/'.join(map(str,ids[aid]['scores'])) and total==ids[aid]['sum8_audit_only']
 ids[aid]['reason_as_submitted']=reason
dc=[]
for aid,question,role,dim,score,reason in wb['逐维度评分理由'].iter_rows(min_row=2,values_only=True):
 i=zh.index(dim);assert score==ids[aid]['scores'][i];dc.append((aid,i))
assert len(dc)==len(set(dc))==108
ties=[]
panel=list(wb['各题得分对照'].iter_rows(min_row=2,values_only=True));assert len(panel)==9
for rn,v in enumerate(panel,2):
 aids=v[2:5];totals=list(v[5:8]);assert totals==[ids[a]['sum8_audit_only'] for a in aids]
 winners=[role for role,total in zip(['甲','乙','丙'],totals) if total==max(totals)]
 if v[8]!='、'.join(winners):
  ties.append(dict(question_number=v[0],cell=f'I{rn}',submitted_winner=v[8],all_tied_roles=winners,all_tied_answer_ids=[aids[['甲','乙','丙'].index(role)] for role in winners]))
summary={}
for name,cases in strata.items():
 groups={}
 for g in ['llm','workflow','agent']:
  rr=[r for r in rows if r['group']==g and r['case_id'] in cases];assert len(rr)==len(cases)
  groups[g]=dict(n_task_episodes=len(rr),dimension_means_0_to_2={d:sum(r['scores'][i] for r in rr)/len(rr) for i,d in enumerate(dims)},completion_label_counts_as_submitted=dict(Counter(r['completion_label_as_submitted'] for r in rr)))
 diffs={a+'_minus_'+b:{d:groups[a]['dimension_means_0_to_2'][d]-groups[b]['dimension_means_0_to_2'][d] for d in dims} for a,b in [('agent','llm'),('workflow','llm'),('agent','workflow')]}
 summary[name]=dict(case_ids=cases,groups=groups,mean_differences_in_scale_points=diffs)
changes=[]
if len(sys.argv)>2:
 old={r['answer_id']:r for r in read(load_workbook(sys.argv[2]))};assert len(old)==27
 for r in rows:
  fields=['scores','completion_label_as_submitted']
  if any(r[f]!=old[r['answer_id']][f] for f in fields):
   changes.append(dict(answer_id=r['answer_id'],before={f:old[r['answer_id']][f] for f in fields},after={f:r[f] for f in fields}))
data=dict(status='USER_CONFIRMED_EXPERT_REVIEW_DESCRIPTIVE_PENDING_CLARIFICATION',source_sha256=sha(src),run_id='mini_real_20261027_061408',rating_provenance='User states an expert reviewed this final file. Reviewer count and independence are not specified in the workbook.',reviewer_count=None,new_ratings_added_by_this_analysis=0,all_cases_are_seen_development_cases=True,unit='One full task episode per group, including both rounds of U1/R1; not 27 independent questions.',score_scale='Four separate ordinal dimensions, 0-2 each. Differences are scale points, not accuracy percentages.',stratified_descriptive_metrics=summary,accepted_primary_completion_rates=None,required_fact_accuracy=None,severe_error_rates=None,primary_metrics_pending_reason='Reviewer clarification is needed for remaining evidence and delivered-answer scope issues. This file lacks protocol per-fact judgments. The serious-error sheet mixes missing deliverables and potentially conclusion-invalidating errors.',scores=rows,changes_from_previous_submission=changes)
save=lambda name,obj:(root/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
save('scores_and_descriptive_summary.json',data)
save('workbook_checks.json',dict(answer_ids_match_all_27_real_positions=True,valid_score_cells=108,reason_rows_match=27,dimension_reason_rows_match=108,total_formulas_checked=27,comparison_totals_match_all_nine_rows=True,tie_labels_omit_other_tied_winners=ties,formula_note='openpyxl does not calculate Excel formulas; totals were independently added from submitted scores. Cached values were not invented.',source_workbook_was_modified=False,previous_workbook_sha256=sha(Path(sys.argv[2])) if len(sys.argv)>2 else None))
with (root/'score_rows_by_actual_group.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['answer_id','case_id','actual_group',*dims,'completion_label_as_submitted','sum8_audit_only'])
 for r in sorted(rows,key=lambda x:(x['case_id'],x['group'])):
  w.writerow([r['answer_id'],r['case_id'],r['group'],*r['scores'],r['completion_label_as_submitted'],r['sum8_audit_only']])
assert sha(src)=='1463b5641298bee179b7c05b9f8b08697afdf856a055c57af7411ea9c995b675'
print(json.dumps(dict(valid_positions=27,valid_score_cells=108,tie_label_issues=ties,changed_score_vectors=sum(c['before']['scores']!=c['after']['scores'] for c in changes),stratified_descriptive_metrics=summary),ensure_ascii=True))

